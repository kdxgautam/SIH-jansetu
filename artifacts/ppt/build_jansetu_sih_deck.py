from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "artifacts" / "ppt"
REFERENCE = Path("/home/hexa/Downloads/Copy of SMART INDIA HACKATHON 2025_20260910_112816_0000.pptx")
PPTX = ROOT / "submission" / "JanSetu_SIH2026_Presentation.pptx"
STACK = OUT / "assets" / "tech-stack.png"
PROOF = OUT / "assets" / "mvp-proof.png"

PAPER = RGBColor(250, 250, 248)
INK = RGBColor(10, 10, 10)
SECONDARY = RGBColor(82, 82, 82)
HELPER = RGBColor(115, 115, 115)
GREY = RGBColor(240, 240, 238)
GREY_2 = RGBColor(212, 212, 210)
BLUE = RGBColor(0, 47, 167)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(29, 92, 69)
FONT = "Noto Sans"
MONO = "Noto Sans Mono"


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def clear_slide(slide):
    for shape in list(slide.shapes):
        remove_shape(shape)


def text_box(slide, x, y, w, h, text, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font=FONT, margin=0.03, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape


def rich_box(slide, x, y, w, h, heading, lines, fill=GREY, accent=BLUE,
             heading_size=18, body_size=15):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.17)
    tf.margin_bottom = Inches(0.12)
    p = tf.paragraphs[0]
    p.space_after = Pt(8)
    r = p.add_run()
    r.text = heading
    r.font.name = MONO
    r.font.size = Pt(heading_size)
    r.font.bold = True
    r.font.color.rgb = accent
    for line in lines:
        p = tf.add_paragraph()
        p.space_after = Pt(4)
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = "- " + line
        r.font.name = FONT
        r.font.size = Pt(body_size)
        r.font.color.rgb = INK
    return shape


def rect(slide, x, y, w, h, fill, line=None, width=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=GREY_2, width=1):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def arrow(slide, x, y, w, h, direction="right", color=BLUE):
    kind = {
        "right": MSO_SHAPE.RIGHT_ARROW,
        "left": MSO_SHAPE.LEFT_ARROW,
        "down": MSO_SHAPE.DOWN_ARROW,
    }[direction]
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def badge(slide, text, x=0.35, y=0.22, w=1.65):
    rect(slide, x, y, w, 0.52, PAPER, BLUE, 1.2)
    text_box(slide, x, y + 0.02, w, 0.44, text, 12, BLUE, True, PP_ALIGN.CENTER, MONO, 0,
             MSO_ANCHOR.MIDDLE)


def add_logo(slide):
    slide.shapes.add_picture(BytesIO(logo_blob), Inches(17.56), Inches(0.11), Inches(2.23), Inches(1.0))


def header(slide, title, page):
    add_logo(slide)
    badge(slide, "BILLUSENA")
    text_box(slide, 2.35, 0.14, 14.7, 0.66, title, 27, INK, True, PP_ALIGN.CENTER)
    line(slide, 2.35, 0.88, 17.1, 0.88, BLUE, 1.4)
    rect(slide, 0, 10.82, 20, 0.43, BLUE)
    text_box(slide, 8.35, 10.87, 3.3, 0.24, "BilluSena · JanSetu", 9, WHITE, False,
             PP_ALIGN.CENTER, MONO, 0, MSO_ANCHOR.MIDDLE)
    text_box(slide, 17.4, 10.87, 1.0, 0.24, f"{page:02} / 06", 9, WHITE, False,
             PP_ALIGN.CENTER, MONO, 0, MSO_ANCHOR.MIDDLE)


def flow_node(slide, x, y, number, role, action, accent=False):
    fill = BLUE if accent else GREY
    color = WHITE if accent else INK
    rect(slide, x, y, 4.0, 1.45, fill)
    text_box(slide, x + 0.22, y + 0.16, 0.55, 0.32, number, 13,
             WHITE if accent else BLUE, True, font=MONO, margin=0)
    text_box(slide, x + 0.92, y + 0.14, 2.78, 0.32, role, 13,
             WHITE if accent else HELPER, True, font=MONO, margin=0)
    text_box(slide, x + 0.22, y + 0.61, 3.55, 0.62, action, 19, color, False, margin=0)


def metric(slide, x, y, w, number, label):
    line(slide, x, y, x + w, y, BLUE, 2)
    text_box(slide, x, y + 0.14, w, 0.58, number, 31, BLUE, False, font=FONT, margin=0)
    text_box(slide, x, y + 0.78, w, 0.55, label, 13, SECONDARY, False, font=FONT, margin=0)


prs = Presentation(str(REFERENCE))
logo_blob = prs.slides[1].part.related_part("rId5").blob
brain = deepcopy(next(
    shape._element for shape in prs.slides[0].shapes
    if shape.left / 914400 > 13 and shape.top / 914400 > 2
))
for slide in prs.slides:
    clear_slide(slide)

prs.core_properties.title = "JanSetu - SIH 26043"
prs.core_properties.subject = "Bilingual societal innovation portal for Jharkhand"
prs.core_properties.author = "BilluSena"


# 01 / Title
slide = prs.slides[0]
add_logo(slide)
slide.shapes._spTree.insert_element_before(deepcopy(brain), "p:extLst")
rect(slide, 0.8, 1.65, 0.16, 8.25, BLUE)
text_box(slide, 1.25, 0.28, 14.7, 0.72, "SMART INDIA HACKATHON 2026", 31, INK, True)
text_box(slide, 1.25, 1.62, 7.8, 0.32, "PS 26043 / SMART EDUCATION / SOFTWARE", 15, BLUE, True,
         font=MONO, margin=0)
text_box(slide, 1.25, 2.14, 9.2, 1.25, "JANSETU", 57, INK, False, margin=0)
text_box(slide, 1.25, 3.35, 9.6, 1.22,
         "Community challenges to validated solutions.", 30, SECONDARY, False, margin=0)
line(slide, 1.25, 4.92, 10.55, 4.92, GREY_2, 1)
text_box(slide, 1.25, 5.25, 10.5, 1.4,
         "A bilingual platform to crowdsource local problems and coordinate Government, universities and industry through one accountable lifecycle.",
         21, INK, False, margin=0)
text_box(slide, 1.25, 7.05, 11.1, 1.62,
         "Organization  Government of Jharkhand\nDepartment   Higher & Technical Education\nTeam         BilluSena\nTeam ID      [Registered Team ID]",
         15, SECONDARY, False, font=MONO, margin=0)
rect(slide, 1.25, 9.35, 3.6, 0.54, BLUE)
text_box(slide, 1.25, 9.38, 3.6, 0.42, "WORKING MVP · DEPLOYED", 13, WHITE, True,
         PP_ALIGN.CENTER, MONO, 0, MSO_ANCHOR.MIDDLE)


# 02 / Proposed solution and complete flow
slide = prs.slides[1]
header(slide, "IDEA TITLE / PROPOSED SOLUTION", 2)
text_box(slide, 0.8, 1.18, 18.4, 0.34, "ONE PLATFORM · ONE ACCOUNTABLE LIFECYCLE", 15, BLUE, True,
         font=MONO, margin=0)
nodes_top = [
    (0.8, "01", "COMMUNITY", "Report need + evidence", False),
    (5.35, "02", "GOVERNMENT", "Review + bilingual public copy", True),
    (9.9, "03", "HEI / UNIVERSITY", "Accept assignment + form team", False),
    (14.45, "04", "PROJECT TEAM", "Proposal + faculty/student roster", False),
]
for i, (x, n, role, action, accent) in enumerate(nodes_top):
    flow_node(slide, x, 1.72, n, role, action, accent)
    if i < 3:
        arrow(slide, x + 4.08, 2.26, 0.37, 0.30, "right")
arrow(slide, 18.55, 3.33, 0.30, 0.54, "down")
nodes_bottom = [
    (14.45, "05", "INDUSTRY", "Mentorship · funding · pilot", False),
    (9.9, "06", "UNIVERSITY", "Build + submit milestone evidence", False),
    (5.35, "07", "GOVERNMENT", "Approve milestones + validate outcome", True),
    (0.8, "08", "COMMUNITY / PUBLIC", "Feedback + approved results", False),
]
for i, (x, n, role, action, accent) in enumerate(nodes_bottom):
    flow_node(slide, x, 4.02, n, role, action, accent)
    if i < 3:
        arrow(slide, x - 0.45, 4.56, 0.37, 0.30, "left")
text_box(slide, 0.8, 5.76, 18.4, 0.28, "CONTROLLED BRANCHES", 13, HELPER, True, font=MONO, margin=0)
branches = [
    ("INFORMATION REQUEST", "Government asks the citizen for missing detail"),
    ("DECLINE / REALLOCATE", "A university decline returns the challenge for allocation"),
    ("REVISION LOOP", "Proposal, milestone or outcome changes return to the university"),
]
for i, (title, body) in enumerate(branches):
    x = 0.8 + i * 6.15
    rect(slide, x, 6.16, 5.75, 1.05, PAPER, GREY_2, 0.8)
    text_box(slide, x + 0.18, 6.32, 1.9, 0.23, title, 11, BLUE, True, font=MONO, margin=0)
    text_box(slide, x + 2.08, 6.24, 3.42, 0.64, body, 13, SECONDARY, False, margin=0)
rect(slide, 0.8, 7.62, 18.4, 2.3, BLUE)
text_box(slide, 1.05, 7.9, 3.0, 0.38, "GEMINI ASSISTS", 18, WHITE, True, font=MONO, margin=0)
text_box(slide, 1.05, 8.48, 12.9, 0.92,
         "Classification  /  priority rationale  /  bilingual drafts  /  duplicate hints  /  university matches  /  outcome evidence gaps",
         16, WHITE, False, margin=0)
rect(slide, 14.65, 8.05, 4.0, 1.42, WHITE)
text_box(slide, 14.92, 8.24, 3.45, 0.32, "HUMANS DECIDE", 15, BLUE, True, font=MONO, margin=0)
text_box(slide, 14.92, 8.76, 3.45, 0.45, "Publish · assign · approve · validate", 13, INK, False, margin=0)


# 03 / Technical approach
slide = prs.slides[2]
header(slide, "TECHNICAL APPROACH", 3)
text_box(slide, 0.8, 1.17, 18.4, 0.44,
         "Browser /api/v1  →  FastAPI business rules  →  PostgreSQL + private evidence  ↔  Gemini advisory services",
         15, BLUE, True, font=MONO, margin=0)
slide.shapes.add_picture(str(STACK), Inches(0.72), Inches(1.68), width=Inches(18.56), height=Inches(7.42))
rect(slide, 0.8, 9.34, 8.85, 1.08, GREY)
text_box(slide, 1.02, 9.53, 1.5, 0.25, "SECURITY", 12, BLUE, True, font=MONO, margin=0)
text_box(slide, 2.52, 9.42, 6.82, 0.62,
         "Argon2 · random DB sessions · HttpOnly cookies · origin checks · role + ownership rules",
         13, SECONDARY, False, margin=0)
rect(slide, 9.9, 9.34, 9.3, 1.08, GREY)
text_box(slide, 10.12, 9.53, 1.65, 0.25, "EVIDENCE", 12, BLUE, True, font=MONO, margin=0)
text_box(slide, 11.8, 9.42, 7.08, 0.62,
         "Private disk storage · signature checks · 5 files × 20 MB · authorised downloads",
         13, SECONDARY, False, margin=0)


# 04 / Feasibility and viability
slide = prs.slides[3]
header(slide, "FEASIBILITY AND VIABILITY", 4)
text_box(slide, 0.8, 1.18, 18.4, 0.36, "BUILDABLE NOW · EXPANDABLE WITH GOVERNMENT ADOPTION", 15, BLUE, True,
         font=MONO, margin=0)
rich_box(slide, 0.8, 1.72, 5.75, 2.72, "01 / TECHNICAL", [
    "Next.js, FastAPI and PostgreSQL are proven production technologies.",
    "Transactional lifecycle and access rules stay in the API.",
    "Structured AI responses are validated before display.",
], body_size=14)
rich_box(slide, 7.12, 1.72, 5.75, 2.72, "02 / ECONOMIC", [
    "Open-source core with usage-based hosting.",
    "Neon, Cloud Run and Vercel scale independently.",
    "Funding commitments remain separate from money received.",
], body_size=14)
rich_box(slide, 13.45, 1.72, 5.75, 2.72, "03 / OPERATIONAL", [
    "Four workspaces follow real stakeholder responsibilities.",
    "Human approval protects government accountability.",
    "Seed and migrations provide a repeatable demonstration.",
], body_size=14)
text_box(slide, 0.8, 4.88, 11.75, 0.30, "RISKS → BUILT-IN RESPONSE", 13, HELPER, True, font=MONO, margin=0)
risks = [
    ("AI unavailable", "Save first; retry transient failures; keep manual review."),
    ("Unsafe public copy", "Minimum lengths + bilingual government approval."),
    ("Access / upload abuse", "Role checks, origin validation and file signatures."),
    ("Ephemeral file storage", "Move evidence to private object storage before rollout."),
]
for i, (risk, response) in enumerate(risks):
    y = 5.34 + i * 1.07
    rect(slide, 0.8, y, 11.75, 0.82, PAPER, GREY_2, 0.8)
    text_box(slide, 1.02, y + 0.19, 2.65, 0.28, risk, 13, INK, True, margin=0)
    text_box(slide, 3.82, y + 0.15, 8.38, 0.42, response, 13, SECONDARY, False, margin=0)
text_box(slide, 13.18, 4.88, 6.02, 0.30, "PROPOSED ROLLOUT", 13, HELPER, True, font=MONO, margin=0)
phases = [
    ("NOW", "Deployed MVP", True),
    ("01", "Pilot districts", False),
    ("02", "Participating HEIs", False),
    ("03", "Statewide network", False),
]
for i, (num, label, accent) in enumerate(phases):
    y = 5.34 + i * 1.08
    rect(slide, 13.18, y, 1.0, 0.82, BLUE if accent else GREY)
    text_box(slide, 13.18, y + 0.16, 1.0, 0.36, num, 12, WHITE if accent else BLUE, True,
             PP_ALIGN.CENTER, MONO, 0, MSO_ANCHOR.MIDDLE)
    text_box(slide, 14.5, y + 0.14, 4.25, 0.44, label, 16, INK, False, margin=0)
    if i < 3:
        line(slide, 13.68, y + 0.82, 13.68, y + 1.08, BLUE, 2)
text_box(slide, 13.18, 9.78, 5.9, 0.42, "Rollout phases are proposed; field adoption is not claimed.",
         12, SECONDARY, False, margin=0)


# 05 / Impact and proof
slide = prs.slides[4]
header(slide, "IMPACT AND BENEFITS", 5)
text_box(slide, 0.8, 1.15, 18.4, 0.30, "MEASUREMENT CHAIN", 13, BLUE, True, font=MONO, margin=0)
chain = ["REPORTED", "ASSIGNED", "DELIVERED", "VALIDATED"]
for i, label in enumerate(chain):
    x = 0.8 + i * 4.62
    rect(slide, x, 1.62, 4.05, 0.72, BLUE if i == 3 else GREY)
    text_box(slide, x, 1.76, 4.05, 0.35, label, 13, WHITE if i == 3 else INK, True,
             PP_ALIGN.CENTER, MONO, 0, MSO_ANCHOR.MIDDLE)
    if i < 3:
        arrow(slide, x + 4.12, 1.83, 0.38, 0.28, "right")
text_box(slide, 0.8, 2.72, 8.25, 0.30, "VALUE FOR EACH PARTICIPANT", 13, HELPER, True, font=MONO, margin=0)
stakeholders = [
    ("01 / CITIZEN", "Bilingual reporting, progress visibility and optional feedback."),
    ("02 / GOVERNMENT", "Review, allocation, approvals, audit history and analytics."),
    ("03 / UNIVERSITY", "Real problem statements, multidisciplinary teams and delivery."),
    ("04 / INDUSTRY", "Reviewed opportunities matched to practical support."),
]
for i, (role, benefit) in enumerate(stakeholders):
    y = 3.14 + i * 1.38
    rect(slide, 0.8, y, 8.25, 1.08, GREY)
    text_box(slide, 1.02, y + 0.17, 2.2, 0.28, role, 12, BLUE, True, font=MONO, margin=0)
    text_box(slide, 3.25, y + 0.13, 5.48, 0.62, benefit, 14, INK, False, margin=0)
text_box(slide, 9.62, 2.72, 9.58, 0.30, "WORKING MVP EVIDENCE", 13, HELPER, True, font=MONO, margin=0)
slide.shapes.add_picture(str(PROOF), Inches(9.62), Inches(3.14), width=Inches(9.58), height=Inches(3.03))
line(slide, 9.62, 6.38, 14.3, 6.38, BLUE, 2)
line(slide, 14.53, 6.38, 19.2, 6.38, BLUE, 2)
text_box(slide, 9.62, 6.53, 4.55, 0.36, "GOVERNMENT REVIEW · HINDI", 11, SECONDARY, True, font=MONO, margin=0)
text_box(slide, 14.53, 6.53, 4.55, 0.36, "INDUSTRY AI MATCHING", 11, SECONDARY, True, font=MONO, margin=0)
text_box(slide, 9.62, 7.06, 9.58, 0.30, "SEEDED DEMO SCALE", 13, HELPER, True, font=MONO, margin=0)
metric(slide, 9.62, 7.52, 2.05, "7", "demo accounts")
metric(slide, 12.08, 7.52, 2.05, "10", "demo challenges")
metric(slide, 14.54, 7.52, 2.05, "5", "institutions")
metric(slide, 17.0, 7.52, 2.05, "6", "lifecycle stages")
text_box(slide, 9.62, 9.22, 9.5, 0.75,
         "Public pages expose approved summaries and aggregate outcomes only. These numbers demonstrate MVP coverage, not measured field impact.",
         12, SECONDARY, False, margin=0)


# 06 / Research and references
slide = prs.slides[5]
header(slide, "RESEARCH AND REFERENCES", 6)
text_box(slide, 0.8, 1.18, 7.45, 0.30, "DEMO + SOURCE", 13, BLUE, True, font=MONO, margin=0)
links = [
    ("FRONTEND", "jansetu-jharkhand.vercel.app"),
    ("API HEALTH", "jansetu-api-595351864949.us-east4.run.app/api/v1/health"),
    ("PRIVATE REPOSITORY", "github.com/kdxgautam/SIH-jansetu"),
]
for i, (label, url) in enumerate(links):
    y = 1.68 + i * 1.62
    rect(slide, 0.8, y, 7.45, 1.25, GREY)
    text_box(slide, 1.02, y + 0.18, 1.55, 0.26, label, 11, BLUE, True, font=MONO, margin=0)
    text_box(slide, 1.02, y + 0.57, 6.8, 0.40, url, 13, INK, False, font=MONO, margin=0)
text_box(slide, 0.8, 6.75, 7.45, 0.30, "PROBLEM CONTEXT", 13, BLUE, True, font=MONO, margin=0)
rich_box(slide, 0.8, 7.18, 7.45, 2.62, "WHY JANSETU", [
    "SIH 26043 asks for crowdsourced societal challenges and collaborative problem solving.",
    "NEP 2020 supports experiential, multidisciplinary and industry-linked learning.",
    "The portal connects those goals through an auditable government workflow.",
], body_size=13)
text_box(slide, 8.72, 1.18, 10.48, 0.30, "PRIMARY TECHNICAL REFERENCES", 13, BLUE, True, font=MONO, margin=0)
refs = [
    ("AI", "Gemini structured output · ai.google.dev/gemini-api/docs/structured-output"),
    ("AUTH", "Vertex ADC · cloud.google.com/vertex-ai/generative-ai/docs/start/quickstart"),
    ("WEB", "Next.js · nextjs.org/docs"),
    ("API", "FastAPI · fastapi.tiangolo.com"),
    ("DATA", "PostgreSQL · postgresql.org/docs   /   Neon · neon.tech/docs"),
    ("CLOUD", "Cloud Run · cloud.google.com/run/docs   /   Vercel · vercel.com/docs"),
]
for i, (label, ref) in enumerate(refs):
    y = 1.67 + i * 1.14
    line(slide, 8.72, y, 19.2, y, GREY_2, 0.8)
    text_box(slide, 8.72, y + 0.18, 1.35, 0.30, label, 11, BLUE, True, font=MONO, margin=0)
    text_box(slide, 10.12, y + 0.13, 8.92, 0.60, ref, 13, INK, False, margin=0)
rect(slide, 8.72, 8.8, 10.48, 1.0, BLUE)
text_box(slide, 8.97, 9.01, 2.05, 0.27, "DEMO STATUS", 12, WHITE, True, font=MONO, margin=0)
text_box(slide, 11.15, 8.94, 7.72, 0.52,
         "Seeded data · no field-impact claims · not an official government service",
         13, WHITE, False, margin=0)


prs.save(str(PPTX))
print(PPTX)
